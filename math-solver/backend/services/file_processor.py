import os
import logging
import asyncio
from typing import Optional, Dict, Any
from concurrent.futures import ThreadPoolExecutor

try:
    import pytesseract
    from PIL import Image
except ImportError:
    pytesseract = None
    Image = None

try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

try:
    from docx import Document
except ImportError:
    Document = None

logger = logging.getLogger(__name__)


class FileProcessor:
    """Process various file types and extract text content"""
    
    # Default configuration
    DEFAULT_MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB
    DEFAULT_ALLOWED_TYPES = ['png', 'jpg', 'jpeg', 'gif', 'bmp', 'pdf', 'docx', 'doc', 'txt']
    DEFAULT_TEXT_ENCODING = 'utf-8'
    TEXT_EXTRACTION_TIMEOUT = 30  # seconds
    
    def __init__(self, max_file_size: Optional[int] = None, allowed_types: Optional[list] = None):
        """
        Initialize file processor
        
        Args:
            max_file_size: Maximum file size in bytes (default 10MB)
            allowed_types: List of allowed file extensions (default see DEFAULT_ALLOWED_TYPES)
        """
        self.max_file_size = max_file_size or self.DEFAULT_MAX_FILE_SIZE
        self.allowed_types = allowed_types or self.DEFAULT_ALLOWED_TYPES
        
        # Set Tesseract path for Windows
        if os.name == 'nt' and pytesseract:
            pytesseract.pytesseract.tesseract_cmd = r'C:\Program Files\Tesseract-OCR\tesseract.exe'
        
        # Thread pool for blocking I/O operations
        self.executor = ThreadPoolExecutor(max_workers=2)
        
        logger.info(
            f"FileProcessor initialized (max_size={self.max_file_size}, "
            f"allowed_types={','.join(self.allowed_types)})"
        )
    
    async def process_file(self, file_path: str, file_type: str) -> Dict[str, Any]:
        """
        Process uploaded file and extract text/content
        
        Args:
            file_path: Path to the file
            file_type: File extension/type
        
        Returns:
            Dict with success status, extracted text, method used, and metadata
        """
        logger.info(f"Processing file: {file_path} (type={file_type})")
        
        try:
            # Validate file exists
            if not os.path.exists(file_path):
                logger.error(f"File not found: {file_path}")
                return self._error_response("File not found", "file_not_found")
            
            # Validate file type
            file_type_lower = file_type.lower()
            if file_type_lower not in self.allowed_types:
                logger.warning(f"Unsupported file type: {file_type}")
                return self._error_response(
                    f"File type .{file_type} not supported. Allowed: {', '.join(self.allowed_types)}",
                    "unsupported_type"
                )
            
            # Validate file size
            file_size = os.path.getsize(file_path)
            if file_size > self.max_file_size:
                logger.warning(f"File too large: {file_size} bytes (max {self.max_file_size})")
                return self._error_response(
                    f"File size {file_size} bytes exceeds maximum {self.max_file_size} bytes",
                    "file_too_large"
                )
            
            # Route to appropriate processor
            if file_type_lower in ['png', 'jpg', 'jpeg', 'gif', 'bmp']:
                return await self._process_image(file_path)
            elif file_type_lower == 'pdf':
                return await self._process_pdf(file_path)
            elif file_type_lower in ['docx', 'doc']:
                return await self._process_docx(file_path)
            elif file_type_lower == 'txt':
                return await self._process_text(file_path)
            else:
                return self._error_response(f"Unsupported file type: {file_type}", "unsupported_type")
        
        except Exception as e:
            logger.exception(f"Unexpected error processing file: {file_path}")
            return self._error_response(f"File processing error: {str(e)}", "processing_error")
    
    async def _process_image(self, file_path: str) -> Dict[str, Any]:
        """Process image file with OCR"""
        logger.info(f"Processing image: {file_path}")
        
        try:
            if not Image or not pytesseract:
                logger.error("PIL or pytesseract not installed")
                return self._error_response(
                    "OCR support not available. Install pillow and pytesseract.",
                    "dependency_missing"
                )
            
            # Check if Tesseract is available (Windows only)
            if os.name == 'nt':
                if not os.path.exists(pytesseract.pytesseract.tesseract_cmd):
                    logger.warning("Tesseract not found, returning placeholder response")
                    return {
                        "success": False,
                        "error": "Tesseract OCR not installed. Install from: https://github.com/UB-Mannheim/tesseract/wiki",
                        "error_type": "dependency_missing",
                        "text": None,
                        "method": "ocr",
                        "fallback": True
                    }
            
            # Run OCR in thread pool
            loop = asyncio.get_event_loop()
            text = await asyncio.wait_for(
                loop.run_in_executor(
                    self.executor,
                    self._ocr_image,
                    file_path
                ),
                timeout=self.TEXT_EXTRACTION_TIMEOUT
            )
            
            if not text or not text.strip():
                logger.warning(f"No text extracted from image: {file_path}")
                return self._error_response(
                    "No text found in image",
                    "no_content"
                )
            
            logger.info(f"Image processed successfully: {len(text)} characters")
            return {
                "success": True,
                "text": text.strip(),
                "method": "ocr",
                "confidence": "medium",
                "content_length": len(text)
            }
        
        except asyncio.TimeoutError:
            logger.error(f"OCR timeout for: {file_path}")
            return self._error_response("OCR processing timeout", "timeout")
        except Exception as e:
            logger.exception(f"OCR processing failed: {file_path}")
            return self._error_response(f"OCR processing failed: {str(e)}", "ocr_error")
    
    def _ocr_image(self, file_path: str) -> str:
        """Blocking OCR operation (runs in executor)"""
        image = Image.open(file_path)
        return pytesseract.image_to_string(image)
    
    async def _process_pdf(self, file_path: str) -> Dict[str, Any]:
        """Process PDF file and extract text"""
        logger.info(f"Processing PDF: {file_path}")
        
        try:
            if not PyPDF2:
                logger.error("PyPDF2 not installed")
                return self._error_response(
                    "PDF support not available. Install PyPDF2.",
                    "dependency_missing"
                )
            
            # Run PDF extraction in thread pool
            loop = asyncio.get_event_loop()
            result = await asyncio.wait_for(
                loop.run_in_executor(
                    self.executor,
                    self._extract_pdf,
                    file_path
                ),
                timeout=self.TEXT_EXTRACTION_TIMEOUT
            )
            
            text, page_count = result
            
            if not text or not text.strip():
                logger.warning(f"No text extracted from PDF: {file_path}")
                return self._error_response(
                    "No text content in PDF",
                    "no_content"
                )
            
            logger.info(f"PDF processed successfully: {page_count} pages, {len(text)} characters")
            return {
                "success": True,
                "text": text.strip(),
                "method": "pdf_extraction",
                "pages": page_count,
                "content_length": len(text)
            }
        
        except asyncio.TimeoutError:
            logger.error(f"PDF extraction timeout: {file_path}")
            return self._error_response("PDF extraction timeout", "timeout")
        except Exception as e:
            logger.exception(f"PDF processing failed: {file_path}")
            return self._error_response(f"PDF processing failed: {str(e)}", "pdf_error")
    
    def _extract_pdf(self, file_path: str) -> tuple:
        """Blocking PDF extraction (runs in executor)"""
        text = ""
        with open(file_path, 'rb') as file:
            try:
                pdf_reader = PyPDF2.PdfReader(file)
                page_count = len(pdf_reader.pages)
                
                for page_num, page in enumerate(pdf_reader.pages):
                    try:
                        text += page.extract_text() + "\n"
                    except Exception as e:
                        logger.warning(f"Failed to extract page {page_num}: {str(e)}")
                        continue
                
                return text, page_count
            except Exception as e:
                logger.exception(f"PDF reader error")
                raise
    
    async def _process_docx(self, file_path: str) -> Dict[str, Any]:
        """Process Word document and extract text"""
        logger.info(f"Processing DOCX: {file_path}")
        
        try:
            if not Document:
                logger.error("python-docx not installed")
                return self._error_response(
                    "DOCX support not available. Install python-docx.",
                    "dependency_missing"
                )
            
            # Run DOCX extraction in thread pool
            loop = asyncio.get_event_loop()
            result = await asyncio.wait_for(
                loop.run_in_executor(
                    self.executor,
                    self._extract_docx,
                    file_path
                ),
                timeout=self.TEXT_EXTRACTION_TIMEOUT
            )
            
            text, para_count = result
            
            if not text or not text.strip():
                logger.warning(f"No text extracted from DOCX: {file_path}")
                return self._error_response(
                    "No text content in document",
                    "no_content"
                )
            
            logger.info(f"DOCX processed successfully: {para_count} paragraphs, {len(text)} characters")
            return {
                "success": True,
                "text": text.strip(),
                "method": "docx_extraction",
                "paragraphs": para_count,
                "content_length": len(text)
            }
        
        except asyncio.TimeoutError:
            logger.error(f"DOCX extraction timeout: {file_path}")
            return self._error_response("Document extraction timeout", "timeout")
        except Exception as e:
            logger.exception(f"DOCX processing failed: {file_path}")
            return self._error_response(f"Document processing failed: {str(e)}", "docx_error")
    
    def _extract_docx(self, file_path: str) -> tuple:
        """Blocking DOCX extraction (runs in executor)"""
        doc = Document(file_path)
        text = ""
        
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        
        return text, len(doc.paragraphs)
    
    async def _process_text(self, file_path: str) -> Dict[str, Any]:
        """Process plain text file"""
        logger.info(f"Processing text file: {file_path}")
        
        try:
            # Run text read in thread pool
            loop = asyncio.get_event_loop()
            text = await asyncio.wait_for(
                loop.run_in_executor(
                    self.executor,
                    self._read_text_file,
                    file_path
                ),
                timeout=self.TEXT_EXTRACTION_TIMEOUT
            )
            
            if not text or not text.strip():
                logger.warning(f"Empty text file: {file_path}")
                return self._error_response("File is empty", "no_content")
            
            logger.info(f"Text file processed successfully: {len(text)} characters")
            return {
                "success": True,
                "text": text.strip(),
                "method": "text_read",
                "content_length": len(text)
            }
        
        except asyncio.TimeoutError:
            logger.error(f"Text read timeout: {file_path}")
            return self._error_response("File read timeout", "timeout")
        except Exception as e:
            logger.exception(f"Text processing failed: {file_path}")
            return self._error_response(f"Text processing failed: {str(e)}", "text_error")
    
    def _read_text_file(self, file_path: str) -> str:
        """Blocking text file read (runs in executor)"""
        with open(file_path, 'r', encoding=self.DEFAULT_TEXT_ENCODING) as file:
            return file.read()
    
    def get_file_info(self, file_path: str) -> Dict[str, Any]:
        """Get basic file information"""
        try:
            if not os.path.exists(file_path):
                return {"error": "File not found"}
            
            file_size = os.path.getsize(file_path)
            file_name = os.path.basename(file_path)
            file_extension = file_name.split('.')[-1].lower() if '.' in file_name else ''
            
            return {
                "success": True,
                "name": file_name,
                "size": file_size,
                "extension": file_extension,
                "size_mb": round(file_size / (1024 * 1024), 2),
                "supported": file_extension in self.allowed_types
            }
        except Exception as e:
            logger.exception(f"Error getting file info: {file_path}")
            return {"success": False, "error": str(e)}
    
    def validate_file(self, file_size: int, file_type: str) -> Dict[str, Any]:
        """Validate file before processing"""
        file_type_lower = file_type.lower()
        
        if file_size > self.max_file_size:
            return {
                "valid": False,
                "error": f"File size {file_size} exceeds maximum {self.max_file_size}",
                "error_type": "file_too_large"
            }
        
        if file_type_lower not in self.allowed_types:
            return {
                "valid": False,
                "error": f"File type .{file_type} not supported. Allowed: {', '.join(self.allowed_types)}",
                "error_type": "unsupported_type"
            }
        
        return {"valid": True}
    
    def _error_response(self, error_msg: str, error_type: str) -> Dict[str, Any]:
        """Create standardized error response"""
        return {
            "success": False,
            "error": error_msg,
            "error_type": error_type,
            "text": None
        }
    
    def cleanup(self):
        """Cleanup resources"""
        self.executor.shutdown(wait=True)